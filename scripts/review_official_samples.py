"""Read official country/cycle sampling tables; no pupil estimates or approvals.

The reviewed table/page selections live in OFFICIAL_SAMPLE_TABLES.json. Extracted
numbers retain source hashes and exact row/page locators. A missing row, a region
with a national name, or an ambiguous numeric layout is an error, never zero.
"""
import argparse
import csv
import hashlib
import json
import re
import logging
from pathlib import Path
from html.parser import HTMLParser


def sha(p):
    with p.open('rb') as f:
        return hashlib.file_digest(f, 'sha256').hexdigest()


def normalize(x):
    # A grade-five row cannot stand in for the retained Norwegian grade-four
    # population. Footnote numbers elsewhere are merely typography.
    s=re.sub(r'\(\s*3\s*\)','gradethree',str(x).lower())
    s=re.sub(r'\(\s*4\s*\)','gradefour',s)
    s=re.sub(r'\(\s*5\s*\)','gradefive',s)
    return re.sub('[^a-z]', '', s)


class Tables(HTMLParser):
    def __init__(self):
        super().__init__(); self.tables=[]; self.table=None; self.row=None; self.cell=None
    def handle_starttag(self, tag, attrs):
        if tag=='table': self.table=[]
        elif tag=='tr' and self.table is not None: self.row=[]
        elif tag in ('td','th') and self.row is not None: self.cell=[]
    def handle_data(self, text):
        if self.cell is not None: self.cell.append(text)
    def handle_endtag(self, tag):
        if tag in ('td','th') and self.cell is not None:
            self.row.append(re.sub(r'\s+',' ',''.join(self.cell)).strip());self.cell=None
        elif tag=='tr' and self.row is not None:
            self.table.append(self.row);self.row=None
        elif tag=='table' and self.table is not None:
            self.tables.append(self.table);self.table=None


def numeric(x, proportion=False):
    s=re.sub(r'\s+','',str(x)).replace(',','').replace('%','').replace('\ufffd','')
    if not re.fullmatch(r'-?\d+(?:\.\d+)?',s):return None
    return float(s)*(100 if proportion else 1)


def extract_table(directory,spec,aliases):
    p=directory/spec['file']
    if sha(p)!=spec['sha256']:raise ValueError('Official source changed: '+p.name)
    rows=[]
    if p.suffix=='.html':
        raw=p.read_bytes()
        try:html=raw.decode('utf-8-sig')
        except UnicodeDecodeError:html=raw.decode('cp1252')
        parser=Tables();parser.feed(html)
        for n,row in enumerate(parser.tables[spec['table']],1):rows.append((row,f"table {spec['table']+1}, row {n}"))
    elif p.suffix=='.xlsx':
        import openpyxl
        w=openpyxl.load_workbook(p,read_only=True,data_only=True);s=w[spec['sheet']]
        for n,row in enumerate(s.values,1):rows.append((list(row),f"sheet {s.title}, row {n}"))
        w.close()
    else:
        from pypdf import PdfReader
        logging.getLogger('pypdf').setLevel(logging.ERROR)
        # Reuse extracted layout only when bound to these exact PDF bytes.
        cache=p.with_suffix('.review-pages.json')
        old=json.loads(cache.read_text(encoding='utf-8')) if cache.exists() else {}
        pages=old.get('pages',{}) if old.get('source_sha256')==spec['sha256'] else {}
        if isinstance(pages,list):pages={str(i+1):t for i,t in enumerate(pages)}
        absent=[pg for pg in spec['pages'] if str(pg) not in pages]
        if absent:
            pdf=PdfReader(p)
            for pg in absent:pages[str(pg)]=pdf.pages[pg-1].extract_text(extraction_mode='layout') or ''
            cache.write_text(json.dumps({'source_sha256':spec['sha256'],'pages':pages}),encoding='utf-8')
        for pg in spec['pages']:
            text=pages[str(pg)]
            if spec.get('start_marker') and re.search(spec['start_marker'],text,re.I):
                text=re.split(spec['start_marker'],text,flags=re.I,maxsplit=1)[1]
            if spec.get('end_marker'):text=re.split(spec['end_marker'],text,flags=re.I,maxsplit=1)[0]
            lines=text.splitlines()
            for n,line in enumerate(lines,1):
                for pattern,count in spec.get('continuation_lines',{}).items():
                    if re.search(pattern,line):line+=' '+' '.join(lines[n:n+count])
                row=re.split(r'\s{2,}',line.strip())
                rows.append((row,f'PDF page {pg}, table row line {n}'))
    answer={}
    for row,locator in rows:
        ci=spec.get('country_column')
        if ci is None and p.suffix=='.pdf':
            # A wrapped label such as 'Northern | Ireland' must never match
            # the national Irish row. Join label cells before numeric columns.
            for start in range(min(3,len(row))):
                if start and not all(re.fullmatch(r'[\d\Wab]*',str(v)) for v in row[:start]):continue
                for end in range(start+1,min(start+4,len(row))+1):
                    key=normalize(' '.join(str(v) for v in row[start:end]))
                    if key in aliases and (end==len(row) or numeric(row[end]) is not None or re.match(r'^\d',str(row[end]))):
                        row=row[:start]+[' '.join(str(v) for v in row[start:end])]+row[end:];ci=start;break
                if ci is not None:break
        elif ci is None:ci=next((i for i in range(min(3,len(row))) if normalize(row[i]) in aliases),None)
        if ci is None or len(row)<=ci:continue
        label=row[ci];key=normalize(label)
        if key not in aliases:continue
        c=aliases[key]
        if spec.get('only_countries') and c not in spec['only_countries']:continue
        if spec.get('exclude_countries') and c in spec['exclude_countries']:continue
        if spec.get('percent_columns'):
            # PDF typography sometimes separates 100% into three cells and
            # leaves the population-comment column empty. Percent signs are
            # the reviewed column boundaries, not whitespace cell counts.
            suffix=' '.join(str(v) for v in row[ci+1:])
            percentages=re.findall(r'(?<![\w.])\d(?:[\d ]*\d)?(?:\s*\.\s*\d(?:[\d ]*\d)?)?\s*%',suffix)
            if not percentages:continue  # prose mentioning a country
            row=row[:ci+1]+percentages
        elif p.suffix=='.pdf' and numeric(row[ci+1] if len(row)>ci+1 else '') is None:
            continue  # prose, not a numeric table row
        if spec.get('numeric_column_count') and len(row)-ci-1!=spec['numeric_column_count']:
            continue  # explicitly selected part of a split official table
        values={}
        for name,index in spec['fields'].items():
            index+=ci
            value=numeric(row[index],spec.get('proportions',False) or name in spec.get('field_proportions',[])) if len(row)>index else None
            if value is None and name not in spec.get('official_unreported',{}).get(c,[]):
                raise ValueError(f'{p.name} {locator} {label}: invalid {name}: {row}')
            values[name]=value
        if c in answer:raise ValueError(f'Ambiguous national row {c} in {p.name}')
        flags=';'.join(str(v) for v in row[:ci] if v not in ('',None))
        answer[c]=dict(values,source_file=p.name,source_sha256=spec['sha256'],
            source_url=spec['url'],locator=locator,official_label=str(label),official_flags=flags)
    return answer


def main():
    a=argparse.ArgumentParser();a.add_argument('--repo',type=Path,required=True)
    a.add_argument('--root',type=Path,required=True);a.add_argument('--output',type=Path,required=True)
    args=a.parse_args();directory=args.root/'scientific_review_20261019/official'
    config_path=args.repo/'06_pipeline/review_20261013/OFFICIAL_SAMPLE_TABLES.json'
    config=json.loads(config_path.read_text(encoding='utf-8'))
    for evidence in config.get('quality_source_review',[])+config.get('mode_evidence',[]):
        if 'file' in evidence and sha(directory/evidence['file'])!=evidence['sha256']:
            raise ValueError('Changed reviewed quality/mode source')
    aliases={normalize(label):c for c,labels in config['country_aliases'].items() for label in labels}
    extracted={};errors=[]
    for spec in config['tables']:
        try:extracted[spec['id']]=extract_table(directory,spec,aliases)
        except (ValueError,KeyError,IndexError) as e:errors.append(str(e))
    source=args.root/'international/SOURCE_RUN_MANIFEST.csv'
    with source.open(encoding='utf-8-sig',newline='') as f:contexts=list(csv.DictReader(f))
    result=[];missing=[];dependencies=[]
    for context in contexts:
        c=context['country'];survey=context['survey'];year=context['year'];key=f'{survey}_{year}'
        row=dict(country=c,survey=survey,year=year)
        policy=config['cycles'][key]
        row.update(policy.get('common',{}))
        for id in policy.get('country_table_overrides',{}).get(c,policy['tables']):
            item=extracted.get(id,{}).get(c)
            if item is None:missing.append((c,key,id));continue
            for k,v in item.items():
                if k.startswith('source_') or k in ('locator','official_label','official_flags'):continue
                if k in row and row[k]!=v:raise ValueError(f'Conflicting {k}: {c} {key}')
                row[k]=v
            dependencies.append(dict(country=c,survey=survey,year=year,table_id=id,**item))
        row['source_context_sha256']=context['raw_sha256']
        row['source_population_note']=context.get('population_note','')
        row['assessment_mode']=policy.get('mode','')
        if c in policy.get('digital',[]):row['assessment_mode']='DIGITAL_MAIN_SAMPLE_WITH_OFFICIAL_TREND_LINK'
        if c in policy.get('paper',[]):row['assessment_mode']='PAPER_MAIN_SAMPLE'
        row['assessment_mode']=policy.get('country_mode',{}).get(c,row['assessment_mode'])
        row['population_qualification']=policy.get('country_notes',{}).get(c,'')
        unreported=policy.get('reviewed_unreported',{}).get(c,[])
        for name in unreported:row[name]=None
        row['official_unreported_fields']=';'.join(unreported)
        notes=[]
        if row.get('overall_exclusion_pct') is not None and row['overall_exclusion_pct']>5:notes.append('EXCLUSION_ABOVE_FIVE_PERCENT')
        if row.get('school_before_pct') is not None and row['school_before_pct']<85:notes.append('SCHOOL_RESPONSE_BELOW_85_BEFORE_REPLACEMENT')
        if row.get('student_response_pct',100)<(80 if survey=='PISA' else 85):notes.append('STUDENT_RESPONSE_BELOW_STUDY_TARGET')
        if row.get('population_coverage_pct') is not None and row['population_coverage_pct']<95:notes.append('TARGET_COVERAGE_BELOW_95_PERCENT')
        if survey!='PISA' and row.get('overall_after_pct') is not None and row['overall_after_pct']<75:notes.append('COMBINED_RESPONSE_BELOW_75_PERCENT')
        if any(v is None for v in row.values()):notes.append('OFFICIAL_TABLE_CONTAINS_UNREPORTED_VALUE_NOT_IMPUTED')
        row['quality_notes']=';'.join(notes)
        row['interpretation']='Quality qualifications accompany estimates; response thresholds alone are not automatic country admission rules.'
        row['review_status']='OFFICIAL_COUNTRY_CYCLE_DOCUMENTATION_REVIEWED' if not any(m[0]==c and m[1]==key for m in missing) else 'OFFICIAL_TABLE_ROW_REQUIRES_REVIEW'
        result.append(row)
    args.output.mkdir(parents=True,exist_ok=True)
    def write(name,rows):
        fields=list(dict.fromkeys(k for row in rows for k in row))
        with (args.output/name).open('w',encoding='utf-8',newline='') as f:
            w=csv.DictWriter(f,fields,lineterminator='\n');w.writeheader();w.writerows(rows)
    write('COUNTRY_CYCLE_SAMPLE_QUALITY.csv',result);write('OFFICIAL_SAMPLE_TABLE_DEPENDENCIES.csv',dependencies)
    manifest=dict(status='OFFICIAL_COUNTRY_CYCLE_REVIEW_COMPLETE' if not missing and not errors else 'EXTRACTION_REVIEW_REQUIRED',
        contexts=len(contexts),complete=sum(r['review_status']=='OFFICIAL_COUNTRY_CYCLE_DOCUMENTATION_REVIEWED' for r in result),
        documented_unreported_contexts=sum(bool(r['official_unreported_fields']) for r in result),
        missing=missing,errors=errors,config_sha256=sha(config_path),source_manifest_sha256=sha(source),
        code_sha256=sha(Path(__file__)),output_sha256={n:sha(args.output/n) for n in ['COUNTRY_CYCLE_SAMPLE_QUALITY.csv','OFFICIAL_SAMPLE_TABLE_DEPENDENCIES.csv']})
    (args.output/'OFFICIAL_SAMPLE_REVIEW.json').write_text(json.dumps(manifest,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({k:v for k,v in manifest.items() if k not in ('missing','errors')}))
    if errors:print(json.dumps({'errors':errors}))


if __name__=='__main__':main()
